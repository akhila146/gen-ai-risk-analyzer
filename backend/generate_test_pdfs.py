import os
import csv
import docx
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors

os.makedirs("test_pdfs", exist_ok=True)
styles = getSampleStyleSheet()

def build_pdf(filename, title, data_rows):
    base_name = filename.replace(".pdf", "")
    
    # 1. Original PDF Generation Code (.pdf)
    pdf_path = f"test_pdfs/{filename}"
    doc = SimpleDocTemplate(pdf_path, pagesize=letter)
    story = [
        Paragraph(f"<b>{title}</b>", styles['Title']),
        Spacer(1, 15)
    ]
    if data_rows:
        table = Table(data_rows, colWidths=[200, 320])
        table.setStyle(TableStyle([
            ('GRID', (0,0), (-1,-1), 0.5, colors.grey),
            ('BACKGROUND', (0,0), (0,-1), colors.whitesmoke),
            ('PADDING', (0,0), (-1,-1), 6),
        ]))
        story.append(table)
    doc.build(story)
    print(f"Generated test_pdfs/{filename}")

    # 2. Added Word Document (.docx)
    docx_path = f"test_pdfs/{base_name}.docx"
    docx_doc = docx.Document()
    if title:
        docx_doc.add_heading(title, level=1)
    for row in data_rows:
        docx_doc.add_paragraph(f"{row[0]}: {row[1]}")
    docx_doc.save(docx_path)
    print(f"Generated test_pdfs/{base_name}.docx")

    # 3. Added Text File (.txt)
    txt_path = f"test_pdfs/{base_name}.txt"
    with open(txt_path, "w", encoding="utf-8") as f:
        if title:
            f.write(f"{title}\n" + "="*40 + "\n")
        for row in data_rows:
            f.write(f"{row[0]}: {row[1]}\n")
    print(f"Generated test_pdfs/{base_name}.txt")

    # 4. Added CSV File (.csv)
    csv_path = f"test_pdfs/{base_name}.csv"
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["Field", "Value"])
        writer.writerows(data_rows)
    print(f"Generated test_pdfs/{base_name}.csv")

# 1. Low Risk[cite: 10]
build_pdf("1_low_risk.pdf", "Applicant Financial & Credit Profile", [
    ["Applicant ID", "APP-2026-001"],
    ["Applicant Name", "Rahul Sharma"],
    ["Contact Number", "+91 9876543210"],
    ["Company Name", "TCS Ltd"],
    ["Job Role", "Senior Systems Architect"],
    ["Monthly Income / Salary", "165,000 INR"],
    ["Credit / CIBIL Score", "785"],
    ["Existing EMIs", "12,000 INR"],
    ["Existing Liabilities", "25,000 INR"],
    ["Requested Loan Amount", "800,000 INR"],
    ["Loan Purpose", "Home Renovation"],
    ["Own House", "Yes"],
    ["Vehicles", "1 Car (Honda City, Fully Paid)"],
    ["Home Loan Status", "None"]
])

# 2. Medium Risk[cite: 10]
build_pdf("2_medium_risk.pdf", "Applicant Financial & Credit Profile", [
    ["Applicant ID", "APP-2026-002"],
    ["Applicant Name", "Kiran Varma"],
    ["Contact Number", "+91 9845123456"],
    ["Company Name", "Apex Logistics"],
    ["Job Role", "Operations Specialist"],
    ["Monthly Income / Salary", "62,000 INR"],
    ["Credit / CIBIL Score", "665"],
    ["Existing EMIs", "24,000 INR"],
    ["Existing Liabilities", "180,000 INR"],
    ["Requested Loan Amount", "600,000 INR"],
    ["Loan Purpose", "Medical Expenses"],
    ["Own House", "Rented"],
    ["Vehicles", "1 Two-Wheeler"],
    ["Home Loan Status", "Active personal loan present"]
])

# 3. High Risk[cite: 10]
build_pdf("3_high_risk.pdf", "Applicant Financial & Credit Profile", [
    ["Applicant ID", "APP-2026-003"],
    ["Applicant Name", "Vikram Rathore"],
    ["Contact Number", "+91 9700011223"],
    ["Company Name", "Freelance Consultant"],
    ["Job Role", "Consultant"],
    ["Monthly Income / Salary", "35,000 INR"],
    ["Credit / CIBIL Score", "540"],
    ["Existing EMIs", "28,000 INR"],
    ["Existing Liabilities", "520,000 INR"],
    ["Requested Loan Amount", "750,000 INR"],
    ["Loan Purpose", "Debt Consolidation"],
    ["Own House", "No"],
    ["Vehicles", "None"],
    ["Home Loan Status", "Multiple defaults noted"]
])

# 4. Missing Mandatory Fields[cite: 10]
build_pdf("4_missing_fields.pdf", "Applicant Financial & Credit Profile (Incomplete)", [
    ["Applicant ID", "APP-2026-004"],
    ["Applicant Name", "Anitha Reddy"],
    ["Contact Number", "+91 9123456789"],
    ["Company Name", "Startup Labs"],
    ["Job Role", "Marketing Lead"],
    ["Monthly Income / Salary", "Not Provided"],
    ["Credit / CIBIL Score", "Pending Verification"],
    ["Existing EMIs", "None Specified"],
    ["Existing Liabilities", "Unknown"],
    ["Requested Loan Amount", "500,000 INR"],
    ["Loan Purpose", "Education"],
    ["Own House", "Yes"],
    ["Vehicles", "1 Car"],
    ["Home Loan Status", "Unknown"]
])

# 5. Empty PDF[cite: 10]
build_pdf("5_empty.pdf", "", [])

# 6. Duplicate applicant with changed values[cite: 10]
build_pdf("6_duplicate_applicant_update.pdf", "Applicant Financial & Credit Profile", [
    ["Applicant ID", "APP-2026-001"],
    ["Applicant Name", "Rahul Sharma"],
    ["Contact Number", "+91 9876543210"],
    ["Company Name", "TCS Ltd"],
    ["Job Role", "Senior Systems Architect"],
    ["Monthly Income / Salary", "170,000 INR"],
    ["Credit / CIBIL Score", "760"],
    ["Existing EMIs", "15,000 INR"],
    ["Existing Liabilities", "28,000 INR"],
    ["Requested Loan Amount", "900,000 INR"],
    ["Loan Purpose", "Home Extension"],
    ["Own House", "Yes"],
    ["Vehicles", "1 Car (Honda City, Fully Paid)"],
    ["Home Loan Status", "No active home loan"]
])